"""
Extract text lines from uploaded bill files.

Supports:
  - Images (delegates to OCR)
  - PDF
  - DOCX and best-effort DOC
  - XLSX / XLSM / XLS and CSV
  - Plain text files
"""

from __future__ import annotations

import csv
import os
import re
from typing import Any, Dict, List

from services.ocr_service import extract_text_payload


IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".webp", ".bmp", ".tif", ".tiff", ".heic"}
DOCUMENT_EXTENSIONS = {".pdf", ".doc", ".docx", ".txt", ".csv", ".xls", ".xlsx", ".xlsm"}
SUPPORTED_EXTENSIONS = IMAGE_EXTENSIONS | DOCUMENT_EXTENSIONS


def _normalise_lines(lines: List[str]) -> List[str]:
    clean: List[str] = []
    for raw in lines:
        line = re.sub(r"\s+", " ", (raw or "")).strip()
        if line:
            clean.append(line)
    return clean


def _extract_from_pdf(file_path: str) -> List[str]:
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    lines: List[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        lines.extend(text.splitlines())
    return _normalise_lines(lines)


def _extract_from_docx(file_path: str) -> List[str]:
    from docx import Document

    doc = Document(file_path)
    lines: List[str] = []
    for para in doc.paragraphs:
        lines.append(para.text)
    return _normalise_lines(lines)


def _extract_from_doc_binary(file_path: str) -> List[str]:
    # Legacy .doc parsing requires external tooling. Use best-effort printable text.
    with open(file_path, "rb") as f:
        data = f.read().decode("latin-1", errors="ignore")

    candidates = re.findall(r"[A-Za-z0-9][A-Za-z0-9\s.,:/()%$#@&+\-]{4,}", data)
    return _normalise_lines(candidates)


def _extract_from_csv(file_path: str) -> List[str]:
    rows: List[str] = []
    with open(file_path, "r", encoding="utf-8-sig", errors="ignore", newline="") as f:
        reader = csv.reader(f)
        for row in reader:
            if row:
                rows.append(" ".join(str(cell) for cell in row if cell is not None))
    return _normalise_lines(rows)


def _extract_from_xlsx(file_path: str) -> List[str]:
    from openpyxl import load_workbook

    wb = load_workbook(filename=file_path, data_only=True, read_only=True)
    rows: List[str] = []
    try:
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                values = [str(cell) for cell in row if cell is not None and str(cell).strip()]
                if values:
                    rows.append(" ".join(values))
    finally:
        wb.close()
    return _normalise_lines(rows)


def _extract_from_xls(file_path: str) -> List[str]:
    import xlrd

    book = xlrd.open_workbook(file_path)
    rows: List[str] = []
    for sheet in book.sheets():
        for r in range(sheet.nrows):
            values = [str(cell) for cell in sheet.row_values(r) if str(cell).strip()]
            if values:
                rows.append(" ".join(values))
    return _normalise_lines(rows)


def _extract_from_txt(file_path: str) -> List[str]:
    with open(file_path, "r", encoding="utf-8-sig", errors="ignore") as f:
        return _normalise_lines(f.read().splitlines())


def extract_bill_text_lines(file_path: str) -> List[str]:
    payload = extract_bill_text_payload(file_path)
    return payload.get("lines", [])


def extract_bill_text_payload(file_path: str) -> Dict[str, Any]:
    _, ext = os.path.splitext(file_path)
    ext = ext.lower()

    if ext in IMAGE_EXTENSIONS:
        return extract_text_payload(file_path)
    if ext == ".pdf":
        return {"lines": _extract_from_pdf(file_path), "tokens": []}
    if ext == ".docx":
        return {"lines": _extract_from_docx(file_path), "tokens": []}
    if ext == ".doc":
        return {"lines": _extract_from_doc_binary(file_path), "tokens": []}
    if ext == ".csv":
        return {"lines": _extract_from_csv(file_path), "tokens": []}
    if ext in {".xlsx", ".xlsm"}:
        return {"lines": _extract_from_xlsx(file_path), "tokens": []}
    if ext == ".xls":
        return {"lines": _extract_from_xls(file_path), "tokens": []}
    if ext == ".txt":
        return {"lines": _extract_from_txt(file_path), "tokens": []}

    raise ValueError(f"Unsupported file format: {ext or '(no extension)'}")
